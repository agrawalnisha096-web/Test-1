#!/usr/bin/env python3
"""Build the SEO Accountability & Results Tracker workbook."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---- Shared style tokens -------------------------------------------------
FONT = "Arial"
NAVY = "1F3864"
BLUE = "2E5496"
LIGHT = "D9E1F2"
YELLOW = "FFF2CC"      # user-input cells
GREY = "F2F2F2"
GREEN = "C6EFCE"
RED = "FFC7CE"
AMBER = "FFEB9C"
WHITE = "FFFFFF"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

H_TITLE = Font(name=FONT, size=18, bold=True, color=NAVY)
H_SUB = Font(name=FONT, size=11, italic=True, color="595959")
H_SECTION = Font(name=FONT, size=12, bold=True, color=NAVY)
H_HEAD = Font(name=FONT, size=10, bold=True, color=WHITE)
BODY = Font(name=FONT, size=10, color="000000")
BODY_B = Font(name=FONT, size=10, bold=True, color="000000")
NOTE = Font(name=FONT, size=9, italic=True, color="7F7F7F")

HEAD_FILL = PatternFill("solid", fgColor=BLUE)
TITLE_FILL = PatternFill("solid", fgColor=LIGHT)
YEL_FILL = PatternFill("solid", fgColor=YELLOW)
GREY_FILL = PatternFill("solid", fgColor=GREY)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

SITES = ["crownseoagency.com", "badassbacklinks.com"]
PARTNERS = ["Partner 1", "Partner 2", "Partner 3"]

wb = Workbook()


def header_row(ws, row, headers, widths, start_col=1):
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = H_HEAD
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(start_col + i)].width = w
    ws.row_dimensions[row].height = 30


def style_body(ws, r1, r2, c1, c2, input_cols=None):
    input_cols = input_cols or []
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font is None or cell.font.name != FONT:
                cell.font = BODY
            cell.border = BORDER
            cell.alignment = LEFT
            if c in input_cols:
                cell.fill = YEL_FILL


def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    for c in cells:
        dv.add(c)


def title_block(ws, title, subtitle, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    t = ws.cell(row=1, column=1, value=title)
    t.font = H_TITLE
    t.fill = TITLE_FILL
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = H_SUB
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


# =========================================================================
# 1. START HERE
# =========================================================================
ws = wb.active
ws.title = "Start Here"
ws.sheet_view.showGridLines = False
title_block(ws, "SEO Accountability & Results Tracker", "Shared tracker for crownseoagency.com + badassbacklinks.com  •  3 partners", 6)

rows = [
    ("", ""),
    ("HOW THIS WORKBOOK WORKS", "section"),
    ("Every tab is a shared log. Fill in the yellow cells; the grey/white cells calculate themselves.", "body"),
    ("Hold ONE 30-minute check-in per week. Walk the 'Weekly Scorecard' tab together, then update it live.", "body"),
    ("", ""),
    ("THE TABS", "section"),
    ("1. Targets & Goals", "Set monthly targets per site; 'Actual' auto-pulls from the Results Dashboard so you see On track / Close / Behind at a glance."),
    ("2. Weekly Scorecard", "Each partner commits to 1-3 goals per week, then marks Done / Partial / Missed. This is the accountability engine."),
    ("3. Activity Log", "Daily record of what each person actually did — task, site, time spent, proof link."),
    ("4. Outreach Pipeline", "Link-building funnel: prospect > contacted > replied > negotiating > won/lost. Won links graduate to the Backlinks tab."),
    ("5. Backlinks", "Every link built or earned: source, DR, anchor, dofollow/nofollow, who built it, live status."),
    ("6. Keyword Rankings", "Track target keywords: start position vs current position, so movement is visible."),
    ("7. On-Page SEO", "Per-page optimization checklist (title, meta, H1, internal links, alt, schema...) with an auto optimization score."),
    ("8. Content", "Every article/page: target keyword, author, word count, publish status, URL."),
    ("9. Content Distribution", "Where each piece gets shared and repurposed — channel, format, owner, status. One post becomes ten touchpoints."),
    ("10. Results Dashboard", "Monthly outcomes per site — clicks, impressions, live backlinks, top-10 keywords, leads."),
    ("", ""),
    ("LEGEND", "section"),
    ("Yellow cell  =  you type here", "legend_y"),
    ("White / grey cell  =  auto-calculated — do not overwrite", "legend_g"),
    ("", ""),
    ("SETUP (do this once)", "section"),
    ("Rename 'Partner 1/2/3' to your real names — use Find & Replace across the whole workbook.", "body"),
    ("Agree on the weekly check-in day/time and put it here:", "body"),
]

r = 4
for label, kind in rows:
    if kind == "section":
        c = ws.cell(row=r, column=1, value=label)
        c.font = H_SECTION
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 22
    elif kind == "body":
        c = ws.cell(row=r, column=1, value="•  " + label)
        c.font = BODY
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c.alignment = LEFT
        ws.row_dimensions[r].height = 18
    elif kind == "legend_y":
        c = ws.cell(row=r, column=1, value=label)
        c.font = BODY_B
        c.fill = YEL_FILL
        c.border = BORDER
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    elif kind == "legend_g":
        c = ws.cell(row=r, column=1, value=label)
        c.font = BODY_B
        c.fill = GREY_FILL
        c.border = BORDER
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    elif kind == "":
        pass
    else:
        # two-column "tab name : description"
        c1 = ws.cell(row=r, column=1, value=label)
        c1.font = BODY_B
        c1.alignment = LEFT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c2 = ws.cell(row=r, column=3, value=kind)
        c2.font = BODY
        c2.alignment = LEFT
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 30
    r += 1

# editable check-in cell
cell = ws.cell(row=r, column=1, value="e.g. Every Monday, 10:00am on Zoom")
cell.fill = YEL_FILL
cell.font = BODY
cell.border = BORDER
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

for col, w in zip("ABCDEF", [26, 16, 20, 14, 14, 14]):
    ws.column_dimensions[col].width = w

# =========================================================================
# 2. WEEKLY SCORECARD  (accountability)
# =========================================================================
ws = wb.create_sheet("Weekly Scorecard")
ws.sheet_view.showGridLines = False
title_block(ws, "Weekly Scorecard", "Each partner sets 1-3 commitments per week, then grades them at the next check-in.", 8)

headers = ["Week Starting", "Partner", "Site", "Commitment (what you'll get done)",
           "Target / Metric", "Status", "Actual Outcome", "Notes / Blockers"]
widths = [14, 12, 20, 34, 18, 12, 26, 26]
HR = 4
header_row(ws, HR, headers, widths)

examples = [
    ("2026-08-17", "Partner 1", "crownseoagency.com", "Publish 2 blog posts targeting 'local SEO' cluster",
     "2 posts live", "Done", "Both live, indexed in 3 days", "—"),
    ("2026-08-17", "Partner 2", "badassbacklinks.com", "Land 3 guest-post backlinks DR40+",
     "3 dofollow links", "Partial", "2 secured, 1 in draft", "Outreach reply slow"),
    ("2026-08-17", "Partner 3", "Both", "Fix all technical audit errors (Screaming Frog)",
     "0 critical errors", "Missed", "Ran out of time", "Rolls to next week"),
]
r = HR + 1
for ex in examples:
    for i, v in enumerate(ex):
        ws.cell(row=r, column=1 + i, value=v)
    r += 1
# blank input rows
last = r + 46
style_body(ws, HR + 1, last, 1, 8, input_cols=[1, 2, 3, 4, 5, 6, 7, 8])

add_dv(ws, '"crownseoagency.com,badassbacklinks.com,Both"', [f"C{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3,All"', [f"B{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Done,Partial,Missed,In Progress"', [f"F{i}" for i in range(HR + 1, last + 1)])

# conditional-ish coloring via manual fills on examples
fills = {"Done": GREEN, "Partial": AMBER, "Missed": RED}
for i, ex in enumerate(examples):
    fc = ws.cell(row=HR + 1 + i, column=6)
    fc.fill = PatternFill("solid", fgColor=fills[ex[5]])
    fc.font = BODY_B

# Summary block (right side)
sc = 10
ws.cell(row=HR, column=sc, value="COMPLETION RATE").font = H_SECTION
sr = HR + 1
sub = ["Partner", "Done", "Partial", "Missed", "Score %"]
for i, h in enumerate(sub):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 12
rng = f"$B${HR+1}:$B${last}"
srng = f"$F${HR+1}:$F${last}"
for j, p in enumerate(PARTNERS):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=p).font = BODY_B
    ws.cell(row=rr, column=sc + 1,
            value=f'=COUNTIFS({rng},{get_column_letter(sc)}{rr},{srng},"Done")')
    ws.cell(row=rr, column=sc + 2,
            value=f'=COUNTIFS({rng},{get_column_letter(sc)}{rr},{srng},"Partial")')
    ws.cell(row=rr, column=sc + 3,
            value=f'=COUNTIFS({rng},{get_column_letter(sc)}{rr},{srng},"Missed")')
    dc = get_column_letter(sc + 1)
    pc = get_column_letter(sc + 2)
    mc = get_column_letter(sc + 3)
    ws.cell(row=rr, column=sc + 4,
            value=f'=IFERROR(({dc}{rr}+0.5*{pc}{rr})/({dc}{rr}+{pc}{rr}+{mc}{rr}),"")')
    ws.cell(row=rr, column=sc + 4).number_format = "0%"
    for cc in range(sc, sc + 5):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
ws.cell(row=sr + 5, column=sc, value="Score = (Done + ½·Partial) ÷ total commitments").font = NOTE
ws.merge_cells(start_row=sr + 5, start_column=sc, end_row=sr + 5, end_column=sc + 4)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:H{last}"

# =========================================================================
# 3. ACTIVITY LOG
# =========================================================================
ws = wb.create_sheet("Activity Log")
ws.sheet_view.showGridLines = False
title_block(ws, "Activity Log", "Daily record of work done. Log it the day you do it — memory fades fast.", 8)

headers = ["Date", "Partner", "Site", "Category", "Task Description",
           "Time (hrs)", "Status", "Proof / Link"]
widths = [13, 12, 20, 18, 40, 10, 13, 30]
HR = 4
header_row(ws, HR, headers, widths)

acts = [
    ("2026-08-14", "Partner 1", "crownseoagency.com", "Content", "Wrote + published 'Local SEO checklist' post", 3.5, "Complete", "https://..."),
    ("2026-08-14", "Partner 2", "badassbacklinks.com", "Link Building", "Outreach to 15 prospects for guest posts", 2.0, "In Progress", "sheet link"),
    ("2026-08-15", "Partner 3", "crownseoagency.com", "Technical", "Fixed 12 broken internal links", 1.5, "Complete", "audit export"),
]
r = HR + 1
for a in acts:
    for i, v in enumerate(a):
        ws.cell(row=r, column=1 + i, value=v)
    ws.cell(row=r, column=6).number_format = "0.0"
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 8, input_cols=[1, 2, 3, 4, 5, 6, 7, 8])
for i in range(HR + 1, last + 1):
    ws.cell(row=i, column=6).number_format = "0.0"

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"C{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"B{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Keyword Research,Content,Link Building,Technical,On-Page,Outreach,Reporting,Other"',
       [f"D{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Complete,In Progress,Blocked"', [f"G{i}" for i in range(HR + 1, last + 1)])

# summary: hours per partner
sc = 10
ws.cell(row=HR, column=sc, value="HOURS LOGGED").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Partner", "Total hrs", "Tasks"]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 12
brng = f"$B${HR+1}:$B${last}"
hrng = f"$F${HR+1}:$F${last}"
for j, p in enumerate(PARTNERS):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=p).font = BODY_B
    ws.cell(row=rr, column=sc + 1, value=f'=SUMIFS({hrng},{brng},{get_column_letter(sc)}{rr})')
    ws.cell(row=rr, column=sc + 1).number_format = "0.0"
    ws.cell(row=rr, column=sc + 2, value=f'=COUNTIFS({brng},{get_column_letter(sc)}{rr})')
    for cc in range(sc, sc + 3):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
tr = sr + 4
ws.cell(row=tr, column=sc, value="TOTAL").font = BODY_B
ws.cell(row=tr, column=sc + 1, value=f'=SUM({get_column_letter(sc+1)}{sr+1}:{get_column_letter(sc+1)}{sr+3})')
ws.cell(row=tr, column=sc + 1).number_format = "0.0"
ws.cell(row=tr, column=sc + 2, value=f'=SUM({get_column_letter(sc+2)}{sr+1}:{get_column_letter(sc+2)}{sr+3})')
for cc in range(sc, sc + 3):
    ws.cell(row=tr, column=cc).border = BORDER
    ws.cell(row=tr, column=cc).alignment = CENTER
    ws.cell(row=tr, column=cc).fill = GREY_FILL
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:H{last}"

# =========================================================================
# 4. BACKLINKS
# =========================================================================
ws = wb.create_sheet("Backlinks")
ws.sheet_view.showGridLines = False
title_block(ws, "Backlinks Tracker", "Every link built or earned. 'Status = Live' is what actually counts.", 11)

headers = ["Date", "Site", "Target URL (yours)", "Source Domain", "DR / DA",
           "Anchor Text", "Type", "Method", "Built By", "Cost ($)", "Status"]
widths = [12, 18, 26, 24, 9, 22, 12, 16, 12, 10, 12]
HR = 4
header_row(ws, HR, headers, widths)

bl = [
    ("2026-08-10", "badassbacklinks.com", "/services", "nichesite.com", 52, "backlink services", "Dofollow", "Guest Post", "Partner 2", 0, "Live"),
    ("2026-08-12", "crownseoagency.com", "/", "marketingblog.com", 41, "SEO agency", "Dofollow", "Outreach", "Partner 1", 120, "Pending"),
    ("2026-08-13", "crownseoagency.com", "/blog/local-seo", "directory.io", 30, "crownseoagency", "Nofollow", "Directory", "Partner 3", 0, "Live"),
]
r = HR + 1
for b in bl:
    for i, v in enumerate(b):
        ws.cell(row=r, column=1 + i, value=v)
    ws.cell(row=r, column=10).number_format = '$#,##0'
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 11, input_cols=list(range(1, 12)))
for i in range(HR + 1, last + 1):
    ws.cell(row=i, column=10).number_format = '$#,##0'

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"B{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Dofollow,Nofollow"', [f"G{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Guest Post,Outreach,Directory,Niche Edit,HARO/Digital PR,Forum,Social,Other"',
       [f"H{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"I{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Live,Pending,Lost,Rejected"', [f"K{i}" for i in range(HR + 1, last + 1)])

# summary per site
sc = 13
ws.cell(row=HR, column=sc, value="LINK SUMMARY").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Site", "Live", "Dofollow (live)", "Spend $"]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 15
sB = f"$B${HR+1}:$B${last}"
sK = f"$K${HR+1}:$K${last}"
sG = f"$G${HR+1}:$G${last}"
sCost = f"$J${HR+1}:$J${last}"
for j, s in enumerate(SITES):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=s).font = BODY_B
    sitecell = f"{get_column_letter(sc)}{rr}"
    ws.cell(row=rr, column=sc + 1, value=f'=COUNTIFS({sB},{sitecell},{sK},"Live")')
    ws.cell(row=rr, column=sc + 2, value=f'=COUNTIFS({sB},{sitecell},{sK},"Live",{sG},"Dofollow")')
    ws.cell(row=rr, column=sc + 3, value=f'=SUMIFS({sCost},{sB},{sitecell})')
    ws.cell(row=rr, column=sc + 3).number_format = '$#,##0'
    for cc in range(sc, sc + 4):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:K{last}"

# =========================================================================
# 5. KEYWORD RANKINGS
# =========================================================================
ws = wb.create_sheet("Keyword Rankings")
ws.sheet_view.showGridLines = False
title_block(ws, "Keyword Rankings", "Positive 'Change' = you moved UP. Update positions on the same day each week.", 10)

headers = ["Site", "Keyword", "Target URL", "Volume", "Difficulty",
           "Start Pos.", "Current Pos.", "Change", "Date Checked", "Owner"]
widths = [18, 26, 26, 10, 11, 11, 12, 10, 13, 12]
HR = 4
header_row(ws, HR, headers, widths)

kw = [
    ("crownseoagency.com", "local seo agency", "/", 1300, 45, 28, 14, "2026-08-14", "Partner 1"),
    ("badassbacklinks.com", "buy backlinks", "/services", 2400, 68, 55, 39, "2026-08-14", "Partner 2"),
    ("crownseoagency.com", "small business seo", "/blog/local-seo", 880, 38, 40, 22, "2026-08-14", "Partner 3"),
]
r = HR + 1
for k in kw:
    ws.cell(row=r, column=1, value=k[0])
    ws.cell(row=r, column=2, value=k[1])
    ws.cell(row=r, column=3, value=k[2])
    ws.cell(row=r, column=4, value=k[3])
    ws.cell(row=r, column=5, value=k[4])
    ws.cell(row=r, column=6, value=k[5])
    ws.cell(row=r, column=7, value=k[6])
    ws.cell(row=r, column=8, value=f"=IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r})),F{r}-G{r},\"\")")
    ws.cell(row=r, column=9, value=k[7])
    ws.cell(row=r, column=10, value=k[8])
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 10, input_cols=[1, 2, 3, 4, 5, 6, 7, 9, 10])
# change column formula for all rows + not an input
for i in range(HR + 1, last + 1):
    cell = ws.cell(row=i, column=8)
    if cell.value is None:
        cell.value = f"=IF(AND(ISNUMBER(F{i}),ISNUMBER(G{i})),F{i}-G{i},\"\")"
    cell.fill = GREY_FILL
    cell.alignment = CENTER

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"A{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"J{i}" for i in range(HR + 1, last + 1)])

# summary: keywords in top 10
sc = 12
ws.cell(row=HR, column=sc, value="RANKING SUMMARY").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Site", "Tracked", "Top 10", "Top 3", "Avg Pos."]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 11
sA = f"$A${HR+1}:$A${last}"
sCur = f"$G${HR+1}:$G${last}"
for j, s in enumerate(SITES):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=s).font = BODY_B
    sitecell = f"{get_column_letter(sc)}{rr}"
    ws.cell(row=rr, column=sc + 1, value=f'=COUNTIFS({sA},{sitecell},{sCur},">0")')
    ws.cell(row=rr, column=sc + 2, value=f'=COUNTIFS({sA},{sitecell},{sCur},"<=10",{sCur},">0")')
    ws.cell(row=rr, column=sc + 3, value=f'=COUNTIFS({sA},{sitecell},{sCur},"<=3",{sCur},">0")')
    ws.cell(row=rr, column=sc + 4, value=f'=IFERROR(AVERAGEIFS({sCur},{sA},{sitecell},{sCur},">0"),"")')
    ws.cell(row=rr, column=sc + 4).number_format = "0.0"
    for cc in range(sc, sc + 5):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:J{last}"

# =========================================================================
# 6. CONTENT
# =========================================================================
ws = wb.create_sheet("Content")
ws.sheet_view.showGridLines = False
title_block(ws, "Content Tracker", "Every page and post — from idea to published.", 9)

headers = ["Site", "Title", "Target Keyword", "Type", "Word Count",
           "Author", "Status", "Publish Date", "URL"]
widths = [18, 34, 22, 14, 12, 12, 14, 13, 28]
HR = 4
header_row(ws, HR, headers, widths)

ct = [
    ("crownseoagency.com", "The 2026 Local SEO Checklist", "local seo checklist", "Blog Post", 1800, "Partner 1", "Published", "2026-08-14", "https://..."),
    ("badassbacklinks.com", "How to Vet a Backlink Vendor", "buy backlinks safely", "Blog Post", 2200, "Partner 2", "Draft", "", ""),
    ("crownseoagency.com", "Services — SEO Audits", "seo audit service", "Landing Page", 900, "Partner 3", "Idea", "", ""),
]
r = HR + 1
for c in ct:
    for i, v in enumerate(c):
        ws.cell(row=r, column=1 + i, value=v)
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 9, input_cols=list(range(1, 10)))

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"A{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Blog Post,Landing Page,Pillar Page,Guest Post,Product/Service,Other"',
       [f"D{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"F{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Idea,Outline,Draft,Review,Published,On Hold"', [f"G{i}" for i in range(HR + 1, last + 1)])

sc = 11
ws.cell(row=HR, column=sc, value="CONTENT SUMMARY").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Site", "Published", "In Pipeline"]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 14
sA = f"$A${HR+1}:$A${last}"
sStat = f"$G${HR+1}:$G${last}"
for j, s in enumerate(SITES):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=s).font = BODY_B
    sitecell = f"{get_column_letter(sc)}{rr}"
    ws.cell(row=rr, column=sc + 1, value=f'=COUNTIFS({sA},{sitecell},{sStat},"Published")')
    ws.cell(row=rr, column=sc + 2, value=f'=COUNTIFS({sA},{sitecell})-COUNTIFS({sA},{sitecell},{sStat},"Published")')
    for cc in range(sc, sc + 3):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:I{last}"

# =========================================================================
# 7. RESULTS DASHBOARD
# =========================================================================
ws = wb.create_sheet("Results Dashboard")
ws.sheet_view.showGridLines = False
title_block(ws, "Results Dashboard", "Monthly outcomes per site. Pull numbers from Google Search Console / Analytics on the 1st.", 9)

headers = ["Month", "Site", "Organic Clicks", "Impressions", "Avg Position",
           "Live Backlinks", "Keywords Top 10", "Leads / Conversions", "Notes"]
widths = [12, 18, 15, 14, 13, 14, 15, 18, 26]
HR = 4
header_row(ws, HR, headers, widths)

res = [
    ("2026-07", "crownseoagency.com", 420, 18500, 22.4, 34, 6, 8, "Baseline month"),
    ("2026-07", "badassbacklinks.com", 260, 12100, 31.0, 41, 3, 5, "Baseline month"),
    ("2026-08", "crownseoagency.com", 610, 24300, 18.9, 39, 9, 12, "Content push paying off"),
    ("2026-08", "badassbacklinks.com", 340, 15600, 27.2, 47, 5, 7, "Link velocity up"),
]
r = HR + 1
for x in res:
    for i, v in enumerate(x):
        ws.cell(row=r, column=1 + i, value=v)
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=4).number_format = "#,##0"
    ws.cell(row=r, column=5).number_format = "0.0"
    r += 1
last = r + 60
style_body(ws, HR + 1, last, 1, 9, input_cols=list(range(1, 10)))
for i in range(HR + 1, last + 1):
    ws.cell(row=i, column=3).number_format = "#,##0"
    ws.cell(row=i, column=4).number_format = "#,##0"
    ws.cell(row=i, column=5).number_format = "0.0"
add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"B{i}" for i in range(HR + 1, last + 1)])

note = ws.cell(row=last + 2, column=1,
               value="Tip: enter Month as text (2026-08) so it sorts cleanly. Fill one row per site per month.")
note.font = NOTE
ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=9)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:I{last}"

# =========================================================================
# 8. TARGETS & GOALS
# =========================================================================
ws = wb.create_sheet("Targets & Goals")
ws.sheet_view.showGridLines = False
title_block(ws, "Targets & Goals", "Set the target; 'Actual' auto-pulls from the Results Dashboard for the reporting month below.", 7)

ws.cell(row=4, column=1, value="Reporting Month:").font = BODY_B
mcell = ws.cell(row=4, column=2, value="2026-08")
mcell.fill = YEL_FILL
mcell.font = BODY_B
mcell.border = BORDER
mcell.alignment = CENTER
ws.cell(row=4, column=3, value="type as YYYY-MM to match the Results Dashboard").font = NOTE

headers = ["Site", "Metric", "Target", "Actual (this month)", "Status", "Gap", "Owner"]
widths = [20, 20, 12, 18, 14, 10, 12]
HR = 6
header_row(ws, HR, headers, widths)

RD = "'Results Dashboard'"
month_ref = "$B$4"
# metric name -> (Results Dashboard column, higher_is_better, sample target for crown)
metrics = [
    ("Organic Clicks", "C", True, 700),
    ("Impressions", "D", True, 22000),
    ("Avg Position", "E", False, 15),
    ("Live Backlinks", "F", True, 45),
    ("Keywords Top 10", "G", True, 12),
    ("Leads / Conversions", "H", True, 15),
]
r = HR + 1
for si, site in enumerate(SITES):
    for name, col, hib, sample in metrics:
        ws.cell(row=r, column=1, value=site)
        ws.cell(row=r, column=2, value=name)
        if si == 0:  # seed example targets for the first site only
            ws.cell(row=r, column=3, value=sample)
        actual = (f"=SUMIFS({RD}!${col}$5:${col}$70,{RD}!$B$5:$B$70,$A{r},"
                  f"{RD}!$A$5:$A$70,{month_ref})")
        ws.cell(row=r, column=4, value=actual)
        if hib:
            status = (f'=IF(C{r}="","Set target",IF(D{r}>=C{r},"On track",'
                      f'IF(D{r}>=0.8*C{r},"Close","Behind")))')
            gap = f'=IF(C{r}="","",D{r}-C{r})'
        else:  # lower is better (Avg Position)
            status = (f'=IF(C{r}="","Set target",IF(AND(D{r}>0,D{r}<=C{r}),"On track",'
                      f'IF(AND(D{r}>0,D{r}<=1.2*C{r}),"Close","Behind")))')
            gap = f'=IF(C{r}="","",C{r}-D{r})'
        ws.cell(row=r, column=5, value=status)
        ws.cell(row=r, column=6, value=gap)
        r += 1
last = r - 1
style_body(ws, HR + 1, last, 1, 7, input_cols=[3, 7])
for rr in range(HR + 1, last + 1):
    for cc in (3, 4, 5, 6):
        ws.cell(row=rr, column=cc).alignment = CENTER
    for cc in (4, 5, 6):
        ws.cell(row=rr, column=cc).fill = GREY_FILL
ws.cell(row=last + 2, column=1,
        value="On track = met target • Close = within ~20% • Behind = further off. "
              "Avg Position is judged in reverse (lower is better).").font = NOTE
ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=7)
ws.freeze_panes = "A7"

# =========================================================================
# 9. OUTREACH PIPELINE
# =========================================================================
ws = wb.create_sheet("Outreach Pipeline")
ws.sheet_view.showGridLines = False
title_block(ws, "Outreach Pipeline", "The link-building funnel. When Status = Won, add the link to the Backlinks tab.", 11)

headers = ["Date Added", "Site", "Prospect Domain", "DR/DA", "Contact (name / email)",
           "Method", "Status", "Follow-up Date", "Est. Cost ($)", "Owner", "Notes"]
widths = [12, 18, 22, 8, 26, 15, 14, 13, 12, 12, 26]
HR = 4
header_row(ws, HR, headers, widths)

op = [
    ("2026-08-11", "badassbacklinks.com", "authorityblog.com", 58, "editor@authorityblog.com", "Guest Post", "Negotiating", "2026-08-19", 150, "Partner 2", "Wants 2 links/mo deal"),
    ("2026-08-12", "crownseoagency.com", "localmarketing.net", 44, "Jane (contact form)", "Niche Edit", "Replied", "2026-08-18", 90, "Partner 1", "Open to insert"),
    ("2026-08-13", "badassbacklinks.com", "seoroundup.com", 61, "team@seoroundup.com", "Digital PR", "Contacted", "2026-08-20", 0, "Partner 3", "Sent data study pitch"),
]
r = HR + 1
for x in op:
    for i, v in enumerate(x):
        ws.cell(row=r, column=1 + i, value=v)
    ws.cell(row=r, column=9).number_format = '$#,##0'
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 11, input_cols=list(range(1, 12)))
for i in range(HR + 1, last + 1):
    ws.cell(row=i, column=9).number_format = '$#,##0'

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"B{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Guest Post,Niche Edit,Outreach,Directory,Digital PR,HARO,Resource Page,Broken Link,Other"',
       [f"F{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Not Contacted,Contacted,Replied,Negotiating,Won,Lost"',
       [f"G{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"J{i}" for i in range(HR + 1, last + 1)])

# funnel summary
sc = 13
ws.cell(row=HR, column=sc, value="PIPELINE FUNNEL").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Stage", "Count"]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 16
stages = ["Not Contacted", "Contacted", "Replied", "Negotiating", "Won", "Lost"]
statrng = f"$G${HR+1}:$G${last}"
for j, st in enumerate(stages):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=st).font = BODY_B
    ws.cell(row=rr, column=sc + 1, value=f'=COUNTIF({statrng},"{st}")')
    for cc in range(sc, sc + 2):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
wr = sr + 1 + len(stages) + 1
ws.cell(row=wr, column=sc, value="Win rate").font = BODY_B
wcell = get_column_letter(sc + 1)
ws.cell(row=wr, column=sc + 1,
        value=f'=IFERROR({wcell}{sr+5}/({wcell}{sr+5}+{wcell}{sr+6}),"")')
ws.cell(row=wr, column=sc + 1).number_format = "0%"
for cc in range(sc, sc + 2):
    ws.cell(row=wr, column=cc).border = BORDER
    ws.cell(row=wr, column=cc).alignment = CENTER
ws.cell(row=wr, column=sc).fill = GREY_FILL
ws.cell(row=wr, column=sc + 1).fill = GREY_FILL
ws.cell(row=wr + 1, column=sc, value="Win rate = Won / (Won + Lost)").font = NOTE
ws.merge_cells(start_row=wr + 1, start_column=sc, end_row=wr + 1, end_column=sc + 1)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:K{last}"

# =========================================================================
# 10. ON-PAGE SEO
# =========================================================================
ws = wb.create_sheet("On-Page SEO")
ws.sheet_view.showGridLines = False
title_block(ws, "On-Page SEO", "One row per page. Mark each check Yes/No — the Score column tallies it automatically.", 14)

checks = ["Title Optimized", "Meta Description", "H1 w/ Keyword", "Internal Links",
          "Image Alt Text", "Schema Markup", "Keyword in URL"]
headers = ["Site", "Page / URL", "Target Keyword"] + checks + ["Word Count", "Score", "Owner", "Status"]
# columns: A Site, B URL, C KW, D..J = 7 checks, K WordCount, L Score, M Owner, N Status
widths = [17, 24, 20] + [13] * 7 + [11, 9, 11, 13]
HR = 4
header_row(ws, HR, headers, widths)

op2 = [
    ("crownseoagency.com", "/blog/local-seo", "local seo checklist",
     "Yes", "Yes", "Yes", "Yes", "No", "No", "Yes", 1800),
    ("badassbacklinks.com", "/services", "buy backlinks",
     "Yes", "No", "Yes", "Yes", "Yes", "No", "No", 1200),
    ("crownseoagency.com", "/", "local seo agency",
     "Yes", "Yes", "No", "Yes", "Yes", "Yes", "No", 900),
]
first_chk = 4   # column D
last_chk = 10   # column J
r = HR + 1
for x in op2:
    for i, v in enumerate(x):
        ws.cell(row=r, column=1 + i, value=v)
    # Score in column L (12)
    ws.cell(row=r, column=12,
            value=f'=COUNTIF({get_column_letter(first_chk)}{r}:{get_column_letter(last_chk)}{r},"Yes")/{len(checks)}')
    ws.cell(row=r, column=12).number_format = "0%"
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 14,
           input_cols=[1, 2, 3] + list(range(first_chk, last_chk + 1)) + [11, 13, 14])
for i in range(HR + 1, last + 1):
    sc_cell = ws.cell(row=i, column=12)
    if sc_cell.value is None:
        sc_cell.value = (f'=COUNTIF({get_column_letter(first_chk)}{i}:'
                         f'{get_column_letter(last_chk)}{i},"Yes")/{len(checks)}')
    sc_cell.number_format = "0%"
    sc_cell.fill = GREY_FILL
    sc_cell.alignment = CENTER
    for cc in range(first_chk, last_chk + 1):
        ws.cell(row=i, column=cc).alignment = CENTER

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"A{i}" for i in range(HR + 1, last + 1)])
for cc in range(first_chk, last_chk + 1):
    add_dv(ws, '"Yes,No,N/A"', [f"{get_column_letter(cc)}{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"M{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"To Do,In Progress,Optimized"', [f"N{i}" for i in range(HR + 1, last + 1)])
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A{HR}:N{last}"

# =========================================================================
# 11. CONTENT DISTRIBUTION
# =========================================================================
ws = wb.create_sheet("Content Distribution")
ws.sheet_view.showGridLines = False
title_block(ws, "Content Distribution & Repurposing", "Every place a piece gets shared or reshaped. One article should spawn many rows here.", 8)

headers = ["Source Content", "Site", "Channel / Platform", "Format",
           "Owner", "Status", "Date", "Link / URL"]
widths = [30, 18, 20, 20, 12, 13, 12, 28]
HR = 4
header_row(ws, HR, headers, widths)

cd = [
    ("The 2026 Local SEO Checklist", "crownseoagency.com", "LinkedIn", "Social Post", "Partner 1", "Published", "2026-08-15", "https://..."),
    ("The 2026 Local SEO Checklist", "crownseoagency.com", "X (Twitter)", "Thread", "Partner 1", "Scheduled", "2026-08-17", ""),
    ("The 2026 Local SEO Checklist", "crownseoagency.com", "YouTube", "Short Video", "Partner 3", "Planned", "", ""),
    ("How to Vet a Backlink Vendor", "badassbacklinks.com", "Newsletter", "Email", "Partner 2", "Planned", "", ""),
    ("How to Vet a Backlink Vendor", "badassbacklinks.com", "Reddit", "Repurposed Article", "Partner 2", "Planned", "", ""),
]
r = HR + 1
for x in cd:
    for i, v in enumerate(x):
        ws.cell(row=r, column=1 + i, value=v)
    r += 1
last = r + 96
style_body(ws, HR + 1, last, 1, 8, input_cols=list(range(1, 9)))

add_dv(ws, '"crownseoagency.com,badassbacklinks.com"', [f"B{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"LinkedIn,X (Twitter),Facebook,Instagram,Reddit,Quora,Newsletter,Medium,YouTube,TikTok,Pinterest,Other"',
       [f"C{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Social Post,Thread,Carousel,Short Video,Infographic,Email,Slide Deck,Podcast,Guest Post,Repurposed Article,Other"',
       [f"D{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Partner 1,Partner 2,Partner 3"', [f"E{i}" for i in range(HR + 1, last + 1)])
add_dv(ws, '"Planned,Scheduled,Published"', [f"F{i}" for i in range(HR + 1, last + 1)])

# channel summary
sc = 10
ws.cell(row=HR, column=sc, value="PUBLISHED BY CHANNEL").font = H_SECTION
sr = HR + 1
for i, h in enumerate(["Channel", "Published"]):
    c = ws.cell(row=sr, column=sc + i, value=h)
    c.font = H_HEAD
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(sc + i)].width = 16
chans = ["LinkedIn", "X (Twitter)", "Facebook", "Instagram", "Reddit", "Newsletter", "Medium", "YouTube"]
crng = f"$C${HR+1}:$C${last}"
strng = f"$F${HR+1}:$F${last}"
for j, ch in enumerate(chans):
    rr = sr + 1 + j
    ws.cell(row=rr, column=sc, value=ch).font = BODY_B
    ws.cell(row=rr, column=sc + 1, value=f'=COUNTIFS({crng},"{ch}",{strng},"Published")')
    for cc in range(sc, sc + 2):
        ws.cell(row=rr, column=cc).border = BORDER
        ws.cell(row=rr, column=cc).alignment = CENTER
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A{HR}:H{last}"

# =========================================================================
# Order the tabs logically, then save
# =========================================================================
desired = ["Start Here", "Targets & Goals", "Weekly Scorecard", "Activity Log",
           "Outreach Pipeline", "Backlinks", "Keyword Rankings", "On-Page SEO",
           "Content", "Content Distribution", "Results Dashboard"]
wb._sheets.sort(key=lambda s: desired.index(s.title))

wb.save("/home/user/Test-1/SEO-Tracker.xlsx")
print("saved", [s.title for s in wb._sheets])
