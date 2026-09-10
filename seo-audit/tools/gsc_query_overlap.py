#!/usr/bin/env python3
"""
GSC query-overlap validator for the white-label consolidation/redirect map.

For every MERGE/REDIRECT row in Trillet-WhiteLabel-Redirect-Map.xlsx, this pulls
the page-level query list for BOTH the source URL and its target from the Search
Console API, then measures how much their ranking queries overlap. High overlap
= safe to merge; low overlap = the two pages serve different intent, keep both.

Run this LOCALLY (it needs your Google credentials). It writes
gsc_overlap_results.csv next to the map.

--------------------------------------------------------------------------------
SETUP
  pip install google-api-python-client google-auth google-auth-oauthlib openpyxl

  Service account (recommended):
    - Download the SA JSON key; add the SA email as a user in Search Console.
    - export GSC_SA_JSON=/path/to/service_account.json

  OAuth (alternative):
    - Download OAuth desktop client_secret.json.
    - export GSC_OAUTH_CLIENT=/path/to/client_secret.json
    - First run opens a browser; a token is cached to gsc_token.json.

  Property (must match exactly):
    export GSC_PROPERTY="sc-domain:trillet.ai"      # or https://www.trillet.ai/
--------------------------------------------------------------------------------
"""
import os, sys, csv, time

MAP_XLSX   = os.path.join(os.path.dirname(__file__), "..", "Trillet-WhiteLabel-Redirect-Map.xlsx")
OUT_CSV    = os.path.join(os.path.dirname(__file__), "..", "gsc_overlap_results.csv")
PROPERTY   = os.environ.get("GSC_PROPERTY", "sc-domain:trillet.ai")
MONTHS     = 16
TOPN       = 100          # top queries per page (by impressions) to compare
SITE_ROOT  = "https://trillet.ai"   # used to build absolute URLs from paths
SCOPES     = ["https://www.googleapis.com/auth/webmasters.readonly"]

# ---------- auth ----------
def service():
    from googleapiclient.discovery import build
    sa = os.environ.get("GSC_SA_JSON")
    if sa:
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_file(sa, scopes=SCOPES)
        return build("searchconsole", "v1", credentials=creds, cache_discovery=False)
    client = os.environ.get("GSC_OAUTH_CLIENT")
    if client:
        from google_auth_oauthlib.flow import InstalledAppFlow
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
        tok = os.path.join(os.path.dirname(__file__), "gsc_token.json")
        creds = None
        if os.path.exists(tok):
            creds = Credentials.from_authorized_user_file(tok, SCOPES)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                creds = InstalledAppFlow.from_client_secrets_file(client, SCOPES).run_local_server(port=0)
            open(tok, "w").write(creds.to_json())
        return build("searchconsole", "v1", credentials=creds, cache_discovery=False)
    sys.exit("Set GSC_SA_JSON or GSC_OAUTH_CLIENT (see header).")

def date_range(months):
    import datetime
    end = datetime.date.today() - datetime.timedelta(days=3)   # GSC lag
    start = end - datetime.timedelta(days=int(months*30.4))
    return start.isoformat(), end.isoformat()

def page_queries(svc, page_url, start, end, topn=TOPN):
    """Return {query: impressions} for one page, top N by impressions."""
    body = {
        "startDate": start, "endDate": end,
        "dimensions": ["query"],
        "dimensionFilterGroups": [{"filters": [
            {"dimension": "page", "operator": "equals", "expression": page_url}]}],
        "rowLimit": topn, "dataState": "final",
    }
    for attempt in range(4):
        try:
            resp = svc.searchanalytics().query(siteUrl=PROPERTY, body=body).execute()
            return {r["keys"][0]: r["impressions"] for r in resp.get("rows", [])}
        except Exception as e:
            if attempt == 3: raise
            time.sleep(2 ** attempt)
    return {}

def abs_url(path):
    return path if path.startswith("http") else SITE_ROOT + path

def load_pairs():
    import openpyxl
    wb = openpyxl.load_workbook(MAP_XLSX)
    ws = wb["Redirect map"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    pairs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        action = row[ci["Action"]]
        src    = row[ci["URL (path)"]]
        tgt    = row[ci["Redirect / merge target"]]
        if action in ("MERGE", "REDIRECT") and tgt:
            pairs.append((src, tgt, action))
    return pairs

def main():
    svc = service()
    start, end = date_range(MONTHS)
    print(f"Property: {PROPERTY}  |  {start} to {end}")
    pairs = load_pairs()
    print(f"{len(pairs)} MERGE/REDIRECT pairs to validate\n")
    cache = {}
    def q(path):
        u = abs_url(path)
        if u not in cache: cache[u] = page_queries(svc, u, start, end)
        return cache[u]
    rows = []
    for src, tgt, action in pairs:
        sq, tq = q(src), q(tgt)
        ss, ts = set(sq), set(tq)
        inter = ss & ts
        # % of the SOURCE's queries that the TARGET also ranks for
        pct_src = (len(inter) / len(ss) * 100) if ss else 0
        jacc = (len(inter) / len(ss | ts) * 100) if (ss or ts) else 0
        if not ss:
            verdict = "SOURCE HAS NO QUERIES — safe to redirect (dead page)"
        elif pct_src >= 40:
            verdict = "HIGH overlap — MERGE confirmed"
        elif pct_src >= 20:
            verdict = "MEDIUM — review the shared queries below"
        else:
            verdict = "LOW overlap — KEEP BOTH (different intent)"
        shared = ", ".join(sorted(inter, key=lambda k: -sq.get(k, 0))[:8])
        rows.append([action, src, tgt, len(ss), len(ts), len(inter),
                     round(pct_src, 1), round(jacc, 1), verdict, shared])
        print(f"[{verdict.split(' —')[0]:>6}] {pct_src:5.1f}% src-overlap | {src}  ->  {tgt}")
    with open(OUT_CSV, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Action","Source","Target","Src queries","Tgt queries",
                    "Shared","% of source shared","Jaccard %","Verdict","Top shared queries"])
        w.writerows(rows)
    print(f"\nWrote {OUT_CSV}")

if __name__ == "__main__":
    main()
